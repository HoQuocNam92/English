from copy import deepcopy
from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

SOURCE = r"C:\Users\quocnam\Downloads\NHOM_NGHIENCUU_DT04 - 2.docx"
OUTPUT = r"D:\workspaces\Projects\English\NHOM_NGHIENCUU_DT04 - Cap nhat Chuong 2 va 3.docx"

doc = Document(SOURCE)
body = doc._element.body

def paragraph_text(el):
    return ''.join(el.itertext()).strip() if el.tag == qn('w:p') else ''

children = list(body)
start = next(i for i, el in enumerate(children) if paragraph_text(el).upper().startswith('CHƯƠNG 2'))
end = next(i for i, el in enumerate(children) if paragraph_text(el).upper().startswith('CHƯƠNG 4'))
for el in children[start:end]:
    body.remove(el)
chapter4 = list(body)[start]

tmp = Document()

def title(text, level=1):
    p = tmp.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    return p

def para(text, bold_lead=None):
    p = tmp.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p

def bullets(items):
    for item in items:
        p = tmp.add_paragraph(style='List Bullet')
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)

def table(headers, rows, widths=None):
    t = tmp.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tmp.add_paragraph()
    return t

title('CHƯƠNG 2: PHÂN TÍCH HỆ THỐNG')
title('2.1. Giới thiệu', 2)
para('Chương này phân tích yêu cầu của hệ thống học tiếng Anh chuyên ngành Công nghệ thông tin dựa trên phạm vi triển khai thực tế. Hệ thống phục vụ ba nhóm người dùng là Học viên, Giảng viên và Quản trị viên; cung cấp Website cho cả ba vai trò và ứng dụng di động cho Học viên. Các nền tảng dùng chung Backend API và cơ sở dữ liệu PostgreSQL để bảo đảm dữ liệu tài khoản, nội dung, bài thi, kết quả và tiến độ được đồng bộ.')

title('2.2. Phân tích yêu cầu chi tiết', 2)
title('2.2.1. Yêu cầu chức năng', 3)
functional = [
('FR01', 'Tài khoản và xác thực', 'Đăng ký, đăng nhập, đăng xuất, đổi mật khẩu, quên mật khẩu bằng OTP, đặt lại mật khẩu và đăng nhập Google trên Website hoặc ứng dụng di động.'),
('FR02', 'Phân quyền', 'Quản lý vai trò, quyền hạn, gán vai trò và kiểm soát truy cập cho Học viên, Giảng viên và Quản trị viên.'),
('FR03', 'Hồ sơ học tập', 'Quản lý trình độ tiếng Anh, lĩnh vực CNTT quan tâm, mục tiêu nghề nghiệp, mục tiêu chứng chỉ và chỉ tiêu học tập của Học viên.'),
('FR04', 'Placement Test', 'Cung cấp bài kiểm tra đầu vào, chấm điểm tại Backend, xác định cấp độ và cập nhật kết quả vào hồ sơ Học viên.'),
('FR05', 'Danh mục nền tảng', 'Quản lý lĩnh vực CNTT, ba cấp độ Beginner, Intermediate, Advanced, mục tiêu nghề nghiệp và chứng chỉ.'),
('FR06', 'Bài học và từ vựng', 'Quản lý bài học, phần nội dung, từ vựng, ví dụ, nguồn tham khảo và liên kết bài học với từ vựng hoặc chứng chỉ.'),
('FR07', 'Technical Reading Lab và Flashcard', 'Học tài liệu kỹ thuật, cập nhật phần trăm hoàn thành và ôn từ vựng bằng flashcard theo bài học.'),
('FR08', 'Ngân hàng câu hỏi', 'Tạo, cập nhật, xóa, tìm kiếm và phân loại câu hỏi theo loại, chủ đề, lĩnh vực, cấp độ và chứng chỉ.'),
('FR09', 'Khảo thí', 'Thiết lập đề, chọn câu hỏi, làm bài, nộp bài, chấm điểm tự động, lưu snapshot, xem kết quả và giải thích đáp án.'),
('FR10', 'Tiến độ và gợi ý', 'Theo dõi mức hoàn thành theo bài học, lĩnh vực và chứng chỉ; tổng hợp kết quả và đề xuất nội dung phù hợp.'),
('FR11', 'Lộ trình học', 'Tạo và lưu lộ trình dựa trên Placement Test, mục tiêu nghề nghiệp, lĩnh vực và thời gian học dự kiến; theo dõi các mô-đun trong lộ trình.'),
('FR12', 'Kế hoạch học tập', 'Tạo, cập nhật, hoàn thành và xóa kế hoạch học theo ngày; thời lượng trong kế hoạch là thời lượng dự kiến.'),
('FR13', 'AI Tutor RAG', 'Hỏi đáp theo học liệu đã duyệt, sửa ngữ pháp, dịch nội dung, giải thích từ vựng, trích dẫn nguồn và từ chối trả lời khi thiếu bằng chứng.'),
('FR14', 'Hỗ trợ học bằng AI', 'Lưu từ vựng từ chatbot, ghi chú cá nhân, đánh giá câu trả lời, ghi nhận lỗi và tạo quiz từ lỗi hoặc lịch sử hội thoại.'),
('FR15', 'Cộng đồng', 'Đăng bài, bình luận, bình chọn, xem chủ đề nổi bật; Quản trị viên có thể khóa hoặc xóa nội dung vi phạm.'),
('FR16', 'Thông báo', 'Gửi và nhận thông báo cá nhân hoặc toàn hệ thống, đếm chưa đọc và đánh dấu đã đọc.'),
('FR17', 'Gói PRO và thanh toán', 'Hiển thị gói 1, 3, 6, 12 tháng; tạo đơn, sinh mã QR, nhận webhook SePay, kích hoạt hoặc gia hạn gói và xem lịch sử giao dịch.'),
('FR18', 'Landing Page và banner', 'Hiển thị trang giới thiệu, bảng giá, chatbot công khai; quản lý banner, CTA, hình ảnh, vị trí và thời gian hiển thị.'),
('FR19', 'Báo cáo', 'Thống kê người dùng, nội dung, kết quả thi, tiến độ, lĩnh vực, cấp độ và mục tiêu nghề nghiệp.'),
('FR20', 'Cài đặt và tệp', 'Chuyển đổi Việt–Anh, sáng–tối và tải ảnh đại diện, bài học hoặc banner qua Cloudinary.'),
]
table(['Mã', 'Nhóm chức năng', 'Yêu cầu'], functional)

title('2.2.2. Yêu cầu phi chức năng', 3)
nonfunctional = [
('NFR01', 'Dễ sử dụng', 'Giao diện thống nhất, có trạng thái tải, lỗi và rỗng; phù hợp trên Website và thiết bị di động.'),
('NFR02', 'Bảo mật', 'Mật khẩu được băm; API dùng JWT, refresh token, kiểm tra vai trò và quyền; dữ liệu đầu vào được xác thực.'),
('NFR03', 'Toàn vẹn dữ liệu', 'Các thao tác nhiều bước như nộp bài và thanh toán phải bảo đảm tính nhất quán; sử dụng khóa, ràng buộc và transaction phù hợp.'),
('NFR04', 'Hiệu năng', 'Phân trang danh sách, lưu số liệu tổng hợp và sử dụng chỉ mục cho các trường thường xuyên tìm kiếm hoặc lọc.'),
('NFR05', 'Khả năng mở rộng', 'Có thể bổ sung lĩnh vực, bài học, chứng chỉ, câu hỏi và nhà cung cấp AI mà không thay đổi kiến trúc tổng thể.'),
('NFR06', 'Độ tin cậy AI', 'Câu trả lời RAG phải gắn với học liệu đã duyệt và trích dẫn nguồn; hệ thống không tự suy đoán khi bằng chứng không đạt ngưỡng.'),
('NFR07', 'Khả dụng', 'Website, mobile và API dùng chung dữ liệu; lỗi từ dịch vụ ngoài phải được thông báo và không làm hỏng dữ liệu nghiệp vụ.'),
]
table(['Mã', 'Thuộc tính', 'Yêu cầu'], nonfunctional)

title('2.3. Mô hình nghiệp vụ', 2)
title('2.3.1. Các tác nhân', 3)
table(['Tác nhân', 'Vai trò và phạm vi'], [
('Học viên', 'Quản lý hồ sơ, làm Placement Test, học bài, flashcard, đọc tài liệu, luyện tập, làm đề, xem kết quả và tiến độ, dùng AI Tutor, quản lý lộ trình, kế hoạch, cộng đồng, thông báo và gói PRO.'),
('Giảng viên', 'Biên soạn bài học, từ vựng, câu hỏi và đề thi; quản lý nội dung chứng chỉ; xem kết quả và tiến độ Học viên theo quyền được cấp.'),
('Quản trị viên', 'Quản lý tài khoản, vai trò, quyền, danh mục, nội dung, báo cáo, AI RAG, cộng đồng, thông báo, banner và hoạt động hệ thống.'),
('Google OAuth', 'Cung cấp danh tính Google cho luồng đăng nhập.'),
('Dịch vụ AI và embedding', 'Sinh nội dung hỗ trợ và vector hóa học liệu cho tìm kiếm ngữ nghĩa.'),
('SePay', 'Cung cấp QR và thông báo giao dịch thanh toán qua webhook.'),
('Cloudinary', 'Lưu trữ ảnh đại diện, ảnh bài học và ảnh banner.'),
])

title('2.3.2. Các quy trình nghiệp vụ chính', 3)
processes = [
('Đăng ký và thiết lập hồ sơ', 'Học viên đăng ký hoặc đăng nhập Google, hoàn thành hồ sơ, chọn lĩnh vực và mục tiêu.'),
('Kiểm tra đầu vào và tạo lộ trình', 'Học viên làm Placement Test; hệ thống chấm điểm, cập nhật cấp độ và tạo lộ trình từ hồ sơ cùng thời gian học dự kiến.'),
('Biên soạn và xuất bản nội dung', 'Giảng viên tạo bài học, phần nội dung, từ vựng và liên kết; nội dung chỉ hiển thị cho Học viên sau khi được xuất bản.'),
('Học và theo dõi tiến độ', 'Học viên học bài, đọc tài liệu hoặc dùng flashcard; hệ thống cập nhật trạng thái và phần trăm hoàn thành.'),
('Thi và chấm điểm', 'Học viên bắt đầu lượt thi, hệ thống lưu snapshot, nhận đáp án, chấm điểm, lưu kết quả và cung cấp giải thích.'),
('Học cùng AI Tutor', 'Câu hỏi được truy xuất trong kho học liệu; AI tạo câu trả lời từ ngữ cảnh, lưu tin nhắn và trích dẫn, đồng thời cho phép lưu từ hoặc đánh giá phản hồi.'),
('Sinh quiz từ lỗi', 'Hệ thống tổng hợp lỗi đã ghi nhận và lịch sử hội thoại để tạo câu hỏi luyện tập phù hợp.'),
('Cộng đồng và kiểm duyệt', 'Học viên đăng bài, bình luận hoặc bình chọn; Quản trị viên khóa và xóa nội dung không phù hợp.'),
('Thanh toán gói PRO', 'Học viên tạo đơn, thanh toán QR; webhook xác nhận giao dịch và hệ thống kích hoạt hoặc gia hạn subscription.'),
('Báo cáo', 'Hệ thống tổng hợp dữ liệu người dùng, nội dung, kết quả và tiến độ để Giảng viên hoặc Quản trị viên theo dõi.'),
]
table(['Quy trình', 'Mô tả'], processes)

title('2.3.3. Danh sách Use Case hệ thống', 3)
usecases = [
('Học viên', 'Đăng ký; đăng nhập; quên mật khẩu; quản lý hồ sơ; Placement Test; xem và học bài; flashcard; đọc tài liệu; luyện tập; làm đề; xem kết quả; xem tiến độ; tạo lộ trình; quản lý kế hoạch; AI Tutor; lưu từ AI; đánh giá AI; quiz từ lỗi; cộng đồng; thông báo; thanh toán PRO.'),
('Giảng viên', 'Đăng nhập; quản lý bài học; quản lý từ vựng; quản lý câu hỏi; thiết lập đề thi; quản lý nội dung chứng chỉ; xem kết quả và tiến độ.'),
('Quản trị viên', 'Quản lý người dùng; quản lý vai trò và quyền; quản lý danh mục; quản lý báo cáo; quản lý RAG; quản lý cộng đồng; quản lý thông báo; quản lý banner.'),
]
table(['Tác nhân', 'Use Case'], usecases)
para('Quan hệ include được sử dụng cho các bước bắt buộc, chẳng hạn Làm đề bao gồm Nộp bài và Chấm điểm tự động. Quan hệ extend được dùng cho các thao tác tùy chọn như Xem giải thích đáp án hoặc Đánh giá câu trả lời AI. Các dịch vụ Google OAuth, AI, SePay và Cloudinary được mô hình hóa là tác nhân ngoài khi chúng trao đổi dữ liệu trực tiếp với hệ thống.')

title('2.4. Mô hình dữ liệu mức phân tích', 2)
title('2.4.1. Các nhóm thực thể', 3)
analysis_entities = [
('Tài khoản và phân quyền', 'User, UserDetail, Role, Permission, AuthenticationToken'),
('Hồ sơ và danh mục', 'LearnerProfile, Domain, Level, CareerGoal, Certificate'),
('Nội dung học tập', 'Lesson, LessonSection, Vocabulary, VocabularyExample, CertificationContent'),
('Khảo thí', 'Question, QuestionOption, Exam, ExamAttempt, AttemptAnswer'),
('Tiến độ và cá nhân hóa', 'LearningProgress, ProgressSummary, Recommendation, LearningPath, LearningPathModule, LearningPlanItem'),
('AI Tutor RAG', 'AIConversation, AIMessage, LearningError, SavedVocabulary, KnowledgeSource, KnowledgeChunk, Citation, Feedback, UsageQuota'),
('Cộng đồng và thông báo', 'DiscussionPost, DiscussionComment, DiscussionVote, Notification'),
('Thương mại và Landing Page', 'PaymentOrder, UserSubscription, LandingBanner'),
]
table(['Nhóm', 'Thực thể mức phân tích'], analysis_entities)
para('Mô hình mức phân tích không bao gồm Smart Review, phiên phân tích thời gian học, chuỗi ngày học, bảng xếp hạng, EXP, huy hiệu, luyện phỏng vấn AI hoặc luyện viết AI vì các chức năng này không thuộc phạm vi triển khai hiện tại.')

title('2.5. Kết chương', 2)
para('Chương 2 đã xác định phạm vi chức năng và dữ liệu phù hợp với sản phẩm hiện tại. Hệ thống tập trung vào nội dung tiếng Anh CNTT, khảo thí, tiến độ, lộ trình, AI Tutor có căn cứ nguồn và các chức năng hỗ trợ vận hành. Kết quả phân tích là cơ sở để thiết kế kiến trúc dữ liệu PostgreSQL và các ràng buộc ở Chương 3.')

title('CHƯƠNG 3: THIẾT KẾ HỆ THỐNG')
title('3.1. Giới thiệu', 2)
para('Chương này trình bày thiết kế cơ sở dữ liệu đang được triển khai bằng PostgreSQL và Prisma ORM. Website Next.js, ứng dụng React Native và Backend NestJS sử dụng chung nguồn dữ liệu. Sau khi loại bỏ các chức năng ngoài phạm vi, cơ sở dữ liệu thực tế có 57 bảng vật lý trong schema public, bao gồm bảng nghiệp vụ, bảng hỗ trợ RAG và bảng lịch sử migration.')

title('3.2. Nguyên tắc thiết kế', 2)
bullets([
'Sử dụng UUID làm khóa chính cho các thực thể nghiệp vụ chính.',
'Chuẩn hóa quan hệ nhiều - nhiều bằng bảng liên kết.',
'Dùng khóa ngoại, unique, enum và index để bảo vệ tính toàn vẹn.',
'Lưu snapshot đề và câu hỏi tại thời điểm Học viên bắt đầu làm bài.',
'Dùng JSONB cho metadata linh hoạt, không thay thế các quan hệ cốt lõi.',
'Dùng vector và chỉ mục HNSW cho tìm kiếm ngữ nghĩa trong RAG.',
'Không lưu dữ liệu Smart Review, streak, huy hiệu hoặc phiên phân tích thời gian học.'
])

groups = [
('3.3.1. Tài khoản và phân quyền', [
('users','Tài khoản xác thực trung tâm.'),('user_details','Thông tin hiển thị và hồ sơ cá nhân.'),('roles','Danh mục vai trò.'),('permissions','Danh mục quyền thao tác.'),('user_roles','Liên kết người dùng với vai trò.'),('role_permissions','Liên kết vai trò với quyền.'),('refresh_tokens','Token làm mới phiên đăng nhập.'),('password_reset_tokens','OTP và yêu cầu đặt lại mật khẩu.')]),
('3.3.2. Danh mục và hồ sơ Học viên', [
('domains','Lĩnh vực CNTT.'),('levels','Ba cấp độ học.'),('career_goals','Mục tiêu nghề nghiệp.'),('career_goal_skills','Kỹ năng cần cho mục tiêu nghề nghiệp.'),('certificates','Danh mục chứng chỉ.'),('certificate_domains','Liên kết chứng chỉ với lĩnh vực.'),('learner_profiles','Hồ sơ học tập của Học viên.'),('learner_profile_domains','Lĩnh vực Học viên quan tâm.'),('learner_profile_career_goals','Mục tiêu nghề nghiệp của Học viên.'),('learner_certificate_goals','Mục tiêu chứng chỉ của Học viên.')]),
('3.3.3. Bài học, từ vựng và chứng chỉ', [
('lessons','Thông tin bài học.'),('lesson_sections','Các phần nội dung trong bài.'),('vocabularies','Từ vựng tiếng Anh CNTT.'),('vocabulary_examples','Ví dụ sử dụng từ vựng.'),('lesson_vocabularies','Liên kết bài học và từ vựng.'),('lesson_certificates','Liên kết bài học và chứng chỉ.'),('certification_contents','Nội dung học theo chứng chỉ.')]),
('3.3.4. Ngân hàng câu hỏi và khảo thí', [
('questions','Ngân hàng câu hỏi.'),('question_options','Phương án trả lời.'),('question_certificates','Liên kết câu hỏi với chứng chỉ.'),('exams','Cấu hình đề thi.'),('exam_questions','Câu hỏi và trọng số trong đề.'),('exam_attempts','Lượt làm bài và snapshot.'),('attempt_answers','Câu trả lời trong lượt thi.'),('attempt_answer_options','Các phương án Học viên đã chọn.')]),
('3.3.5. Tiến độ, gợi ý và lộ trình', [
('learning_progress','Tiến độ theo bài học, lĩnh vực hoặc chứng chỉ.'),('progress_summary_cache','Dữ liệu tổng hợp phục vụ truy xuất nhanh.'),('recommendations','Gợi ý nội dung học.'),('learning_paths','Lộ trình cá nhân của Học viên.'),('learning_path_modules','Các mô-đun trong lộ trình.'),('learning_plan_items','Kế hoạch học theo ngày.')]),
('3.3.6. AI Tutor và RAG', [
('ai_conversations','Phiên hội thoại AI.'),('ai_messages','Tin nhắn người dùng và AI.'),('ai_learning_errors','Lỗi được AI phát hiện và giải thích.'),('ai_saved_vocabulary','Từ hoặc cụm từ lưu từ AI.'),('knowledge_sources','Nguồn học liệu được index.'),('knowledge_chunks','Các đoạn học liệu sau khi chia nhỏ.'),('knowledge_vectors','Embedding dùng cho tìm kiếm ngữ nghĩa.'),('ai_message_citations','Trích dẫn của câu trả lời AI.'),('ai_feedback','Đánh giá câu trả lời AI.'),('ai_usage_daily','Hạn mức và token AI theo ngày.')]),
('3.3.7. Cộng đồng và thông báo', [
('discussion_posts','Bài viết cộng đồng.'),('discussion_comments','Bình luận.'),('discussion_votes','Bình chọn của người dùng.'),('notifications','Thông báo cá nhân hoặc toàn hệ thống.')]),
('3.3.8. Thanh toán và Landing Page', [
('payment_orders','Đơn thanh toán gói PRO.'),('user_subscriptions','Gói và thời hạn sử dụng.'),('landing_banners','Nội dung và cấu hình banner.')]),
('3.3.9. Vận hành', [('_prisma_migrations','Lịch sử migration của Prisma.')]),
]

title('3.3. Thiết kế các nhóm dữ liệu', 2)
count = 0
for heading, rows in groups:
    title(heading, 3)
    count += len(rows)
    table(['STT', 'Tên bảng', 'Vai trò'], [(i + 1, name, desc) for i, (name, desc) in enumerate(rows)])
assert count == 57, count

title('3.4. Các quan hệ chính', 2)
relations = [
('users – roles', 'N–N qua user_roles; vai trò nhận quyền qua role_permissions.'),
('learner_profiles – domains/career_goals', 'N–N qua các bảng liên kết hồ sơ.'),
('lessons – vocabularies/certificates', 'N–N qua lesson_vocabularies và lesson_certificates.'),
('exams – questions', 'N–N qua exam_questions; exam_attempts lưu từng lượt làm.'),
('learning_paths – learning_path_modules', '1–N; mỗi lộ trình gồm các mô-đun có thứ tự.'),
('ai_conversations – ai_messages', '1–N; tin nhắn AI có thể có nhiều citation và feedback.'),
('knowledge_sources – knowledge_chunks', '1–N; mỗi chunk có vector tương ứng phục vụ truy xuất.'),
('discussion_posts – comments/votes', '1–N; vote unique theo cặp bài viết và người dùng.'),
('users – payment_orders/user_subscriptions', 'Một người có nhiều đơn hàng và tối đa một subscription hiện hành.'),
]
table(['Quan hệ', 'Thiết kế'], relations)

title('3.5. Ràng buộc toàn vẹn và nghiệp vụ', 2)
bullets([
'Email người dùng là duy nhất; trạng thái tài khoản và nội dung được giới hạn bằng enum.',
'Mỗi tiến độ là duy nhất theo Học viên, loại tài nguyên và định danh tài nguyên.',
'Một câu trả lời trong lượt thi là duy nhất theo lượt thi và câu hỏi.',
'Đơn thanh toán sử dụng idempotency key; webhook được kiểm tra để tránh xử lý lặp.',
'Từ vựng lưu từ AI là duy nhất theo người dùng, từ và cụm từ.',
'Đánh giá AI là duy nhất theo tin nhắn và người dùng.',
'Chỉ học liệu đã xuất bản được đưa vào kho tri thức RAG.',
'Các thao tác nộp bài, thanh toán và cập nhật nhiều bản ghi phải được xử lý bằng transaction.'
])

title('3.6. Ánh xạ từ phân tích sang thiết kế', 2)
mapping = [
('User', 'users, user_details, roles, permissions, user_roles, role_permissions'),
('LearnerProfile', 'learner_profiles và các bảng liên kết domain, career goal, certificate goal'),
('Lesson và Vocabulary', 'lessons, lesson_sections, vocabularies, vocabulary_examples và các bảng liên kết'),
('Exam', 'questions, question_options, exams, exam_questions, exam_attempts, attempt_answers, attempt_answer_options'),
('LearningProgress', 'learning_progress, progress_summary_cache, recommendations'),
('LearningPath', 'learning_paths, learning_path_modules, learning_plan_items'),
('AI Tutor RAG', 'ai_conversations, ai_messages, ai_learning_errors, ai_saved_vocabulary, knowledge_sources, knowledge_chunks, knowledge_vectors, ai_message_citations, ai_feedback, ai_usage_daily'),
('Community', 'discussion_posts, discussion_comments, discussion_votes'),
('Payment và Landing Page', 'payment_orders, user_subscriptions, landing_banners'),
]
table(['Khái niệm mức phân tích', 'Bảng mức thiết kế'], mapping)

title('3.7. Kết chương', 2)
para('Chương 3 đã cập nhật thiết kế cơ sở dữ liệu theo trạng thái triển khai hiện tại với 57 bảng vật lý. Thiết kế bao phủ tài khoản và RBAC, hồ sơ, nội dung, khảo thí, tiến độ, lộ trình, AI Tutor RAG, cộng đồng, thông báo, thanh toán và Landing Page. Các bảng của Smart Review, gamification, huy hiệu và phân tích thời gian không còn thuộc lược đồ. Thiết kế này là cơ sở cho phần cài đặt thực nghiệm và kiểm thử ở các chương tiếp theo.')

for el in list(tmp._element.body):
    if el.tag != qn('w:sectPr'):
        body.insert(body.index(chapter4), deepcopy(el))

# Keep headings and tables together where possible, and make table headers repeat.
for p in doc.paragraphs:
    if p.style and p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
for t in doc.tables:
    if t.rows:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        if trPr.find(qn('w:tblHeader')) is None:
            trPr.append(OxmlElement('w:tblHeader'))

doc.save(OUTPUT)
print(OUTPUT)
