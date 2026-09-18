import json
import sys
from docx import Document

doc = Document(sys.argv[1])
payload = {
    "paragraphs": [
        {"style": paragraph.style.name if paragraph.style else "", "text": paragraph.text}
        for paragraph in doc.paragraphs if paragraph.text.strip()
    ],
    "tables": [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ],
}
with open(sys.argv[2], "w", encoding="utf-8") as output:
    json.dump(payload, output, ensure_ascii=False, indent=2)
